
import argparse, os, glob, json, numpy as np, pandas as pd
from sklearn.metrics import average_precision_score, precision_recall_curve, fbeta_score, precision_score, recall_score
import tensorflow as tf
from tensorflow.keras.models import load_model

def ensure_dir(p):
    os.makedirs(p, exist_ok=True)

def load_val_npz(path):
    data = np.load(path, allow_pickle=True)
    X = data["X"]
    y = data["y"].astype(int)
    mask = data["mask"] if "mask" in data.files else None
    return X, y, mask

def predict_with_mask_support(model, X, mask):
    # Handle models that expect either 1 input (X) or 2 inputs ([X, mask]).
    try:
        n_inputs = len(model.inputs)
    except Exception:
        # Fallback: assume single input
        n_inputs = 1
    if n_inputs >= 2 and mask is not None:
        try:
            # Try names 'seq' and 'mask' ordering, but Keras doesn't require names to match
            y_prob = model.predict([X, mask], verbose=0).reshape(-1)
            return y_prob
        except Exception:
            pass
    # Single input
    y_prob = model.predict(X, verbose=0).reshape(-1)
    return y_prob

def best_f2_metrics(y_true, y_prob, beta=2.0, num_thresh=200):
    # Sweep thresholds and find the threshold maximizing F-beta
    thresholds = np.linspace(0, 1, num_thresh)
    best = {"F2": 0.0, "thr": 0.5, "recall": 0.0, "precision": 0.0}
    for t in thresholds:
        y_pred = (y_prob >= t).astype(int)
        if y_pred.sum() == 0 and y_true.sum() == 0:
            f2 = 1.0
        else:
            f2 = fbeta_score(y_true, y_pred, beta=beta, zero_division=0)
        if f2 > best["F2"]:
            best["F2"] = f2
            best["thr"] = float(t)
            best["recall"] = recall_score(y_true, y_pred, zero_division=0)
            best["precision"] = precision_score(y_true, y_pred, zero_division=0)
    return best

def extract_config_summary(model):
    """Get a concise summary of the most relevant hyperparameters from the model config."""
    try:
        cfg = model.get_config()
    except Exception:
        return ""
    # Try to detect architecture and key params
    layers = cfg.get("layers") or []
    units = []
    dropouts = []
    conv_kernels = []
    dilations = []
    arch_hint = None
    for l in layers:
        lc = (l.get("class_name") or "").lower()
        conf = l.get("config") or {}
        if "lstm" in lc or "gru" in lc:
            arch_hint = arch_hint or ("bilstm" if "lstm" in lc and "bidirectional" in lc else
                                      "bigru" if "gru" in lc and "bidirectional" in lc else
                                      "lstm" if "lstm" in lc else "gru")
            u = conf.get("units")
            if u is not None:
                units.append(u)
            dr = conf.get("dropout")
            if dr is not None:
                dropouts.append(dr)
        if "conv1d" in lc:
            k = conf.get("kernel_size")
            if k is not None:
                conv_kernels.append(k if isinstance(k, int) else (k[0] if isinstance(k, (list, tuple)) else k))
        if "tcn" in lc or "temporalconvnet" in lc:
            # custom layers may encode dilation in config
            d = conf.get("dilation_rate")
            if d is not None:
                dilations.append(d if isinstance(d, int) else (d[0] if isinstance(d, (list, tuple)) else d))
    # Optimizer LR
    lr = None
    try:
        lr_v = getattr(model.optimizer, "learning_rate", None) or getattr(model.optimizer, "lr", None)
        if lr_v is not None:
            lr = float(tf.keras.backend.get_value(lr_v))
    except Exception:
        pass

    parts = []
    if units:
        parts.append(f"units={units}")
    if dropouts:
        parts.append(f"dropout={round(float(np.mean(dropouts)),3)}")
    if conv_kernels:
        parts.append(f"kernel={conv_kernels}")
    if dilations:
        parts.append(f"dilation={dilations}")
    if lr is not None:
        parts.append(f"lr={lr:g}")
    return "; ".join(parts)

def eval_model(path, X, y, mask):
    model = load_model(path, compile=False)
    y_prob = predict_with_mask_support(model, X, mask)
    pr_auc = float(average_precision_score(y, y_prob))
    f2 = best_f2_metrics(y, y_prob, beta=2.0)
    cfg = extract_config_summary(model)
    return {"PR-AUC (val)": pr_auc, "F2 (val)": f2["F2"], "thr*": f2["thr"],
            "Recall*": f2["recall"], "Precisão*": f2["precision"], "config": cfg}

def main():
    ap = argparse.ArgumentParser(description="Gera Tabela 9.6 — HPO por arquitetura (ganhos vs baseline)")
    ap.add_argument("--val_npz", required=True, help="Ficheiro .npz com X,y,mask (val)")
    ap.add_argument("--baseline_dir", required=True, help="Pasta com baseline .keras (ex: reports/seq_models_quick/artifacts)")
    ap.add_argument("--hpo_dir", required=True, help="Pasta com modelos HPO .keras (ex: models_seq)")
    ap.add_argument("--out_csv", required=True, help="CSV de saída (Tabela 9.6)")
    ap.add_argument("--out_docx", default=None, help="Opcional: DOCX formatado para o Word")
    args = ap.parse_args()

    X, y, mask = load_val_npz(args.val_npz)

    # Mapas de nomes de arquitetura
    archs = ["bilstm", "bigru", "gru", "lstm", "cnn1d", "tcn"]

    rows = []
    for arch in archs:
        # baseline path
        base_candidates = glob.glob(os.path.join(args.baseline_dir, f"{arch}.keras"))
        baseline_path = base_candidates[0] if base_candidates else None

        # HPO candidates
        hpo_candidates = glob.glob(os.path.join(args.hpo_dir, f"top*{arch}*.keras")) + \
                         glob.glob(os.path.join(args.hpo_dir, f"*best*{arch}*.keras")) + \
                         glob.glob(os.path.join(args.hpo_dir, f"*{arch}*_hpo*.keras"))
        best_hpo = None
        best_score = -1.0
        best_eval = None
        for p in hpo_candidates:
            try:
                ev = eval_model(p, X, y, mask)
                if ev["PR-AUC (val)"] > best_score:
                    best_score = ev["PR-AUC (val)"]
                    best_hpo = p
                    best_eval = ev
            except Exception as e:
                # ignore models that fail to load
                pass

        # Evaluate baseline if exists
        baseline_eval = None
        if baseline_path and os.path.exists(baseline_path):
            try:
                baseline_eval = eval_model(baseline_path, X, y, mask)
            except Exception:
                baseline_eval = None

        if baseline_eval is None and best_eval is None:
            continue

        row = {
            "Arquitetura": arch.upper(),
            "Melhor config (HPO)": best_eval["config"] if best_eval else "",
            "PR-AUC (val) baseline": baseline_eval["PR-AUC (val)"] if baseline_eval else "",
            "PR-AUC (val) HPO": best_eval["PR-AUC (val)"] if best_eval else "",
            "ΔPR-AUC": (best_eval["PR-AUC (val)"] - baseline_eval["PR-AUC (val)"]) if (best_eval and baseline_eval) else "",
            "F2 (val) baseline": baseline_eval["F2 (val)"] if baseline_eval else "",
            "F2 (val) HPO": best_eval["F2 (val)"] if best_eval else "",
            "ΔF2": (best_eval["F2 (val)"] - baseline_eval["F2 (val)"]) if (best_eval and baseline_eval) else "",
            "Thr* (HPO)": best_eval["thr*"] if best_eval else "",
            "Recall* (HPO)": best_eval["Recall*"] if best_eval else "",
            "Precisão* (HPO)": best_eval["Precisão*"] if best_eval else "",
            "Modelo HPO": os.path.basename(best_hpo) if best_hpo else "",
            "Modelo baseline": os.path.basename(baseline_path) if baseline_path else ""
        }
        rows.append(row)

    if not rows:
        raise SystemExit("❌ Não encontrei modelos para avaliar. Verifica as pastas baseline_dir e hpo_dir.")

    df = pd.DataFrame(rows)
    # formatação numérica
    for c in ["PR-AUC (val) baseline","PR-AUC (val) HPO","ΔPR-AUC","F2 (val) baseline","F2 (val) HPO","ΔF2",
              "Thr* (HPO)","Recall* (HPO)","Precisão* (HPO)"]:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").map(lambda x: f"{x:.3f}" if pd.notna(x) else "")

    ensure_dir(os.path.dirname(args.out_csv) or ".")
    df.to_csv(args.out_csv, index=False)

    print("✅ Tabela 9.6 (CSV) →", args.out_csv)
    print(df.to_string(index=False))

    # DOCX opcional
    if args.out_docx:
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
            doc = Document()
            style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(11)
            title = doc.add_paragraph(); run = title.add_run("Tabela 9.6 — HPO por arquitetura (melhor configuração e ganhos)"); run.bold = True
            doc.add_paragraph("")
            table = doc.add_table(rows=1+len(df), cols=len(df.columns))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            for j, h in enumerate(df.columns):
                p = table.rows[0].cells[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(h); r.bold = True
            for i, (_, row) in enumerate(df.iterrows(), start=1):
                for j, h in enumerate(df.columns):
                    txt = str(row[h])
                    p = table.rows[i].cells[j].paragraphs[0]
                    # right-align numbers
                    if any(k in h for k in ["PR-AUC","F2","Δ","Thr","Recall","Precisão"]):
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.add_run(txt)
            doc.add_paragraph("")
            note = doc.add_paragraph(); note_run = note.add_run("Métrica em validação; F2 e (Recall/Precisão) medidos no threshold que maximizou F2. "); note_run.italic = True
            doc.save(args.out_docx)
            print("✅ DOCX →", args.out_docx)
        except Exception as e:
            print("⚠️ Não foi possível criar DOCX:", e)

if __name__ == "__main__":
    main()
